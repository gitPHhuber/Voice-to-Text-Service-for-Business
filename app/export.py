"""
Экспорт результатов: DOCX, TXT, JSON, SRT
"""

import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def fmt_time(seconds: float) -> str:
    s = int(seconds)
    h, s = divmod(s, 3600)
    m, s = divmod(s, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def fmt_srt_time(seconds: float) -> str:
    ms = int(round((seconds - int(seconds)) * 1000))
    s = int(seconds)
    h, s = divmod(s, 3600)
    m, s = divmod(s, 60)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def to_json(result, output_path: str):
    """Экспорт в JSON."""
    data = {
        "language": result.language,
        "duration": result.duration,
        "processing_time": result.processing_time,
        "speakers": result.speakers,
        "generated_at": datetime.now().isoformat(),
        "summary": result.summary,
        "topics": result.topics,
        "segments": [
            {
                "start": s.start,
                "end": s.end,
                "speaker": s.speaker,
                "text": s.text,
                "confidence": getattr(s, "confidence", 1.0),
            }
            for s in result.segments
        ],
    }
    Path(output_path).write_text(json.dumps(data, ensure_ascii=False, indent=2), "utf-8")


def to_txt(result, output_path: str):
    """Экспорт в TXT с таймкодами и спикерами."""
    lines = [
        "# Транскрипт",
        f"# Длительность: {fmt_time(result.duration)}",
        f"# Спикеров: {len(result.speakers)}",
        f"# Создано: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
        "",
    ]

    if result.summary:
        lines.append("══ РЕЗЮМЕ ══")
        lines.append(result.summary)
        lines.append("")

    lines.append("══ ТРАНСКРИПТ ══")
    lines.append("")

    for seg in result.segments:
        ts = f"{fmt_time(seg.start)} → {fmt_time(seg.end)}"
        speaker = seg.speaker or "—"
        lines.append(f"[{ts}] {speaker}: {seg.text}")

    Path(output_path).write_text("\n".join(lines), "utf-8")


def to_srt(result, output_path: str):
    """Экспорт в SRT (субтитры)."""
    blocks = []
    for i, seg in enumerate(result.segments, 1):
        ts = f"{fmt_srt_time(seg.start)} --> {fmt_srt_time(seg.end)}"
        speaker_prefix = f"{seg.speaker}: " if seg.speaker else ""
        blocks.append(f"{i}\n{ts}\n{speaker_prefix}{seg.text}\n")
    Path(output_path).write_text("\n".join(blocks), "utf-8")


def regenerate_exports(task_id: str, result, title: str = "Транскрипт"):
    """Перегенерирует все форматы экспорта для транскрипта."""
    from app.config import settings
    out = settings.output_dir / task_id

    # Clean old exports (except status.json)
    for f in out.glob("*"):
        if f.name == "status.json":
            continue
        if f.suffix in (".docx", ".txt", ".json", ".srt"):
            f.unlink()

    base = str(out / task_id)
    to_docx(result, base + ".docx", title)
    to_txt(result, base + ".txt")
    to_json(result, base + ".json")
    to_srt(result, base + ".srt")


def to_docx(result, output_path: str, title: str = "Транскрипт встречи"):
    """Экспорт в форматированный DOCX."""
    doc = Document()

    # Style
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15

    # Title
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Meta
    meta = f"Длительность: {fmt_time(result.duration)}"
    meta += f"  |  Спикеров: {len(result.speakers)}"
    meta += f"  |  Сегментов: {len(result.segments)}"
    meta += f"\nСоздано: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    doc.add_paragraph(meta)

    # Summary
    if result.summary:
        doc.add_heading("Резюме", level=2)
        doc.add_paragraph(result.summary)

    # Topics TOC
    if result.topics:
        doc.add_heading("Содержание", level=2)
        for topic in result.topics:
            ts = fmt_time(topic.get("start", 0))
            p = doc.add_paragraph(f"  [{ts}] {topic.get('title', '')}")

    # Transcript
    doc.add_page_break()
    doc.add_heading("Полный транскрипт", level=2)

    topic_idx = 0
    for i, seg in enumerate(result.segments):
        # Insert topic header
        if topic_idx < len(result.topics) and result.topics[topic_idx].get("start_idx") == i:
            t = result.topics[topic_idx]
            doc.add_heading(t.get("title", ""), level=3)
            topic_idx += 1

        ts = f"{fmt_time(seg.start)} → {fmt_time(seg.end)}"
        p = doc.add_paragraph()

        # Timestamp
        run_ts = p.add_run(f"[{ts}] ")
        run_ts.font.color.rgb = RGBColor(128, 128, 128)
        run_ts.font.size = Pt(9)

        # Speaker
        if seg.speaker:
            run_sp = p.add_run(f"{seg.speaker}: ")
            run_sp.bold = True

        # Text
        p.add_run(seg.text)

    doc.save(output_path)
